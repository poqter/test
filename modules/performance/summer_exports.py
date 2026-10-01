"""Summer result workbook generation.

Extracted from the existing implementation; public facade names are preserved.
"""
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, Border, Side, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.table import Table, TableStyleInfo

from modules.performance.summer_core import (
    TABLE_SEQ,
    mark,
    won,
    safe_table_name,
    autosize_columns_full,
    to_styled,
    format_summary_for_display,
)

def excel_safe_value(value):
    """계산 결과는 유지하면서 Excel 셀에 기록 가능한 값으로만 변환합니다."""
    if value is None:
        return None

    try:
        missing = pd.isna(value)
        if isinstance(missing, (bool, np.bool_)) and missing:
            return None
    except (TypeError, ValueError):
        pass

    if isinstance(value, np.generic):
        return value.item()

    return value

def write_table(ws, df_for_sheet: pd.DataFrame, start_row: int = 1, name_suffix: str = "A"):
    global TABLE_SEQ

    r_idx = start_row

    for r_idx, row in enumerate(dataframe_to_rows(df_for_sheet, index=False, header=True), start_row):
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(
                row=r_idx,
                column=c_idx,
                value=excel_safe_value(value),
            )
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(horizontal="center", vertical="center")

    end_col_letter = ws.cell(row=start_row, column=max(df_for_sheet.shape[1], 1)).column_letter
    last_row = r_idx if df_for_sheet.shape[0] > 0 else start_row

    TABLE_SEQ += 1
    display_name = safe_table_name(f"tbl_{ws.title}_{name_suffix}_{TABLE_SEQ}")

    table = Table(displayName=display_name, ref=f"A{start_row}:{end_col_letter}{last_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    ws.add_table(table)

    # 화면과 동일하게 조정·예외 적용 셀을 강조합니다.
    if "적용 구분" in df_for_sheet.columns:
        headers = {cell.value: cell.column for cell in ws[start_row]}
        target_headers = [
            "원본 쉐어율", "적용 쉐어율", "전체 보험료 역산", "실적보험료",
            "조정 차액", "인정 건수", "썸머율", "적용 구분",
        ]
        for row_num in range(start_row + 1, last_row + 1):
            label = str(ws.cell(row=row_num, column=headers["적용 구분"]).value or "")
            has_share = "쉐어" in label
            has_dental = "치아보험" in label
            fill_color = None
            if has_share and has_dental:
                fill_color = "EEE3FF"
            elif has_dental:
                fill_color = "E3F2FD"
            elif has_share:
                fill_color = "FFF4CC"
            if fill_color:
                for header in target_headers:
                    if header in headers:
                        ws.cell(row=row_num, column=headers[header]).fill = PatternFill("solid", fgColor=fill_color)

    autosize_columns_full(ws, padding=5)

    return last_row

def write_title(ws, row, title):
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = Font(bold=True, size=13)
    cell.alignment = Alignment(horizontal="left", vertical="center")

def write_final_result_block(ws, row, result):
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    fill = PatternFill("solid", fgColor="F2F2F2")

    rows = [
        ["7월 환산업적", won(result["7월"]["환산금액"])],
        ["7월 한화생명 환산업적 합계 5만원 이상", mark(result["7월"]["한화생명5만"])],
        ["7월 한화생명 인정 건수", f"{result['7월']['한화생명인정건수']:g}건"],
        ["7월 한화생명 인정 건수 1건 이상", mark(result["7월"]["한화생명1건"])],
        ["7월 환산업적 50만원 이상", mark(result["7월"]["환산50만"])],
        ["7월 조건 달성", mark(result["7월"]["월달성"])],
        ["8월 환산업적", won(result["8월"]["환산금액"])],
        ["8월 한화생명 환산업적 합계 5만원 이상", mark(result["8월"]["한화생명5만"])],
        ["8월 한화생명 인정 건수", f"{result['8월']['한화생명인정건수']:g}건"],
        ["8월 한화생명 인정 건수 1건 이상", mark(result["8월"]["한화생명1건"])],
        ["8월 환산업적 50만원 이상", mark(result["8월"]["환산50만"])],
        ["8월 조건 달성", mark(result["8월"]["월달성"])],
        ["기본 7월+8월 합산 환산업적", won(result["기본합산환산금액"])],
        ["레디포썸머 보너스율", f"{result['레디포썸머보너스율']:.0f} %"],
        ["레디포썸머 보너스금액", won(result["레디포썸머보너스금액"])],
        ["보너스 반영 최종 환산업적", won(result["합산환산금액"])],
        ["월별 필수조건", mark(result["월별필수조건"])],
        ["금액 기준 등급", result["금액기준등급"]],
        ["최종 인정 등급", result["최종인정등급"]],
    ]

    if result["다음등급"]:
        rows.append([
            f"다음 등급({result['다음등급']})까지 부족금액",
            won(result["다음등급부족금액"]),
        ])
    else:
        rows.append(["최고 등급 달성", "HWARANG"])

    for i, row_data in enumerate(rows, start=row):
        for j, value in enumerate(row_data, start=1):
            cell = ws.cell(row=i, column=j, value=value)
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

            if j == 1:
                cell.fill = fill
                cell.font = Font(bold=True)

    autosize_columns_full(ws, padding=5)

    return row + len(rows)

def build_workbook(
    df_all: pd.DataFrame,
    july_df: pd.DataFrame,
    august_df: pd.DataFrame,
    other_month_df: pd.DataFrame,
    summary: pd.DataFrame,
    result: dict,
    excluded_disp: pd.DataFrame,
    selected_collector: str = "전체",
    review_disp: pd.DataFrame | None = None,
):
    wb = Workbook()

    ws_summary = wb.active
    ws_summary.title = "요약"

    write_title(ws_summary, 1, f"썸머 최종 결과 - {selected_collector}")
    next_row = write_final_result_block(ws_summary, 2, result)

    write_title(ws_summary, next_row + 2, "수금자별 요약")
    next_row = write_table(
        ws_summary,
        format_summary_for_display(summary),
        start_row=next_row + 3,
        name_suffix="SUMMARY",
    )

    write_title(ws_summary, next_row + 2, "상세 내역")
    next_row = write_table(
        ws_summary,
        to_styled(df_all),
        start_row=next_row + 3,
        name_suffix="DETAIL",
    )

    ws_july = wb.create_sheet("7월")
    write_title(ws_july, 1, f"7월 썸머 환산 결과 - {selected_collector}")
    write_table(ws_july, to_styled(july_df), start_row=2, name_suffix="JULY_DETAIL")

    ws_august = wb.create_sheet("8월")
    write_title(ws_august, 1, f"8월 썸머 환산 결과 - {selected_collector}")
    write_table(ws_august, to_styled(august_df), start_row=2, name_suffix="AUGUST_DETAIL")

    if not other_month_df.empty:
        ws_other = wb.create_sheet("7월8월외")
        write_title(ws_other, 1, f"7월/8월 외 계약 - {selected_collector}")
        write_table(ws_other, to_styled(other_month_df), start_row=2, name_suffix="OTHER_MONTH")

    if excluded_disp is not None and not excluded_disp.empty:
        ws_ex = wb.create_sheet("제외계약")
        write_title(ws_ex, 1, f"제외 계약 - {selected_collector}")
        write_table(ws_ex, excluded_disp, start_row=2, name_suffix="EXCLUDED")

    if review_disp is not None and not review_disp.empty:
        ws_review = wb.create_sheet("확인필요계약")
        write_title(ws_review, 1, f"입력값 확인이 필요한 계약 - {selected_collector}")
        write_table(ws_review, review_disp, start_row=2, name_suffix="REVIEW")

    return wb
