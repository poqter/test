"""Commission workbook presentation/export.

Extracted from the existing implementation; public facade names are preserved.
"""
from __future__ import annotations
from modules.shared.runtime_cache import session_export
import io
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from modules.performance.commission_calculator_core import (
    _compact_product_display,
    _collector_label,
)

@session_export("commission-excel-v2", name_arg=False)
def _make_excel(
    contracts: list[dict], payout_rate: float, reference_month: str, excluded: list[dict],
    fallback_collectors: list[str] | None = None,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "수수료 계산"
    total_premium = sum(item["premium"] for item in contracts)
    total_first = sum(item["premium"] * item["first_year_rate"] * payout_rate for item in contracts)
    total_commission = sum(item["premium"] * item["total_rate"] * payout_rate for item in contracts)
    collector_label = _collector_label(contracts, fallback_collectors)
    title = f"{collector_label} 수수료 계산 결과" if collector_label else "수수료 계산 결과"
    ws.append([title])
    ws.append(["수수료표 기준월", reference_month or "확인 필요", "공통 지급율", payout_rate])
    ws.append(["계약 수", len(contracts), "월보험료 합계", total_premium])
    ws.append(["예상 익월수당 합계", round(total_first), "예상 총수당 합계", round(total_commission)])
    ws.append([])
    headers = ["고객명", "증권번호", "보험회사", "상품 및 세부 조건", "월보험료", "모집 정보",
               "익월 수수료율", "총수수료율", "예상 익월수당", "예상 총수당"]
    ws.append(headers)

    for contract in contracts:
        first_rate = contract["first_year_rate"] * payout_rate
        total_rate = contract["total_rate"] * payout_rate
        premium = contract["premium"]
        product_detail = _compact_product_display(contract, contracts)
        share_rate = contract.get("share_rate", 100.0)
        recruiter_type = contract.get("recruiter_type", "")
        recruiting = f"{share_rate:g}%"
        if share_rate < 100 and recruiter_type:
            recruiting += f" · {recruiter_type}"
        ws.append([
            contract.get("customer", ""),
            contract.get("policy_number", ""),
            contract["insurer"],
            product_detail,
            premium,
            recruiting,
            first_rate,
            total_rate,
            round(premium * first_rate),
            round(premium * total_rate),
        ])

    header_fill = PatternFill("solid", fgColor="2563D9")
    ws.merge_cells("A1:J1")
    ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1E3A8A")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    for cell in ws[6]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws["D2"].number_format = "0%"
    for cell in (ws["D3"], ws["A4"], ws["C4"]):
        cell.font = Font(bold=True)
    ws["B3"].number_format = '0"건"'
    for cell in (ws["D3"], ws["B4"], ws["D4"]):
        cell.number_format = '#,##0"원"'
    for row in range(7, ws.max_row + 1):
        ws.cell(row, 5).number_format = '#,##0"원"'
        for col in range(7, 9):
            ws.cell(row, col).number_format = "0.0%"
        for col in range(9, 11):
            ws.cell(row, col).number_format = '#,##0"원"'
        ws.row_dimensions[row].height = 42

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    widths = [15, 22, 18, 64, 18, 20, 18, 18, 21, 21]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A7"
    ws.auto_filter.ref = f"A6:J{ws.max_row}"
    ws.row_dimensions[1].height = 28

    review_ws = wb.create_sheet("검토 제외 계약")
    review_headers = ["고객명", "증권번호", "보험회사", "상품명", "계약상태", "제외 사유"]
    review_ws.append(review_headers)
    for item in excluded:
        review_ws.append([
            item.get("customer", ""), item.get("policy_number", ""), item.get("insurer", ""),
            _compact_product_display({"product": item.get("product", ""), "insurer": item.get("insurer", "")}),
            item.get("status", ""), item.get("reason", ""),
        ])
    for cell in review_ws[1]:
        cell.fill = PatternFill("solid", fgColor="64748B")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    for col, width in enumerate([15, 22, 18, 64, 15, 55], start=1):
        review_ws.column_dimensions[get_column_letter(col)].width = width
    for row in range(2, review_ws.max_row + 1):
        review_ws.row_dimensions[row].height = 36
    for row in review_ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    review_ws.freeze_panes = "A2"
    if review_ws.max_row > 1:
        review_ws.auto_filter.ref = review_ws.dimensions

    output = io.BytesIO()
    for sheet in wb:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    wb.save(output)
    return output.getvalue()
