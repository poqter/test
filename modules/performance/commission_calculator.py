from __future__ import annotations

# Stable facade: UI only; engines/parsers/exports are independently testable.
from modules.performance.commission_calculator_core import (
    APP_VERSION,
    DEFAULT_PAYOUT_RATE,
    FIRST_YEAR_HEADERS,
    TOTAL_HEADERS,
    PRODUCT_HEADERS,
    PRODUCT_FALLBACK_HEADERS,
    ProductRate,
    HoldingContract,
    _normalize,
    _clean_text,
    _number,
    _effective_max_col,
    _is_header,
    _header_positions,
    _condition_header,
    _condition_value,
    _condition_text,
    _source_payout_rate,
    _extract_sheet,
    parse_commission_workbook,
    _to_product_rate,
    _format_rate,
    _format_won,
    OUTPUT_INSURER_PREFIXES,
    _remove_output_insurer_prefix,
    _output_product_name,
    _contract_payment_label,
    _contract_renewal_label,
    _rate_distinguishing_labels,
    _compact_product_display,
    INSURER_ALIASES,
    _standard_insurer,
    _source_type_from_insurer,
    _month_from_filename,
    _date_text,
    _holding_product_name,
    _product_family_name,
    _product_display_parts,
    _condition_display,
    _condition_pairs,
    _short_condition_label,
    _condition_option_label,
    parse_holding_workbook,
    _payment_matches,
    _has_payment_condition,
    _payment_condition_label,
    _payment_threshold,
    _most_specific_payment_candidates,
    _selection_tags,
    _tag_match_summary,
    _filter_by_holding_tags,
    _condition_sort_key,
    _sort_condition_candidates,
    PRODUCT_CATEGORY_TOKENS,
    _product_categories,
    _strip_revision_markers,
    _smart_product_name,
    _structural_signatures,
    _bigrams,
    _name_similarity,
    _hard_product_conflict,
    _rank_products,
    _ranked_product_groups,
    _candidate_products,
    _review_candidate_products,
    _auto_candidate,
    _analyze_product_links_uncached,
    _analyze_product_links,
    _collector_label,
    _commission_download_filename,
    _holding_caption,
    _markdown_text,
    _product_groups,
    _direct_product_names,
)

from modules.performance.commission_calculator_exports import (
    _make_excel,
)


# 전달용 파일: 월별 수수료표 스마트 연결 흐름 적용본 v14

import hashlib
import copy
from modules.shared.runtime_cache import cached, fingerprint, digest_bytes, session_export
import io
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from difflib import SequenceMatcher
from typing import Any

import streamlit as st
from modules.shared.upload_ui import guarded_upload
from modules.shared.ui_components import page_footer, page_header, section_intro

import streamlit.components.v1 as components
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _initialize_state() -> None:
    st.session_state.setdefault("commission_contracts", [])
    from modules.shared.organization import active_profile
    profile = active_profile()
    revision_key = (profile["revision"], profile["effective_from"])
    if st.session_state.get("commission_organization_revision") != revision_key:
        st.session_state["commission_payout_rate"] = float(profile["commission_default_payout_percent"])
        st.session_state["commission_organization_revision"] = revision_key
    st.session_state.setdefault("commission_payout_rate", DEFAULT_PAYOUT_RATE)
    st.session_state.setdefault("commission_edit_index", None)
    st.session_state.setdefault("commission_edit_request", None)
    st.session_state.setdefault("commission_ratebook_signature", "")
    st.session_state.setdefault("commission_import_collectors", [])
    st.session_state.setdefault("commission_import_contract_months", [])


def _reconnect_contract_rates(contracts: list[dict], products: list[ProductRate]) -> tuple[int, int]:
    """새 수수료표에서 기존 상품·세부 조건이 같은 계약의 요율을 다시 연결합니다."""
    updated_count = 0
    unresolved_count = 0
    for contract in contracts:
        same_product = [
            product for product in products
            if product.source_type == contract.get("source_type")
            and product.insurer == contract.get("insurer")
            and _holding_product_name(product.product) == _holding_product_name(contract.get("product", ""))
        ]
        same_condition = [
            product for product in same_product
            if _normalize(_condition_display(product)) == _normalize(contract.get("conditions", "") or "기본 조건")
        ]
        candidates = same_condition or same_product
        unique_rates = {
            (round(product.first_year_rate, 8), round(product.total_rate, 8)) for product in candidates
        }
        if not candidates or len(unique_rates) != 1:
            contract["rate_recheck_required"] = True
            unresolved_count += 1
            continue
        selected = candidates[0]
        contract.update({
            "product": selected.product, "conditions": selected.conditions,
            "first_year_rate": selected.first_year_rate, "total_rate": selected.total_rate,
            "sheet_name": selected.sheet_name, "row_number": selected.row_number,
            "rate_recheck_required": False,
        })
        updated_count += 1
    return updated_count, unresolved_count


def _contract_data(holding: dict, product: ProductRate, recruiter_type: str = "") -> dict:
    return {
        "customer": holding.get("customer", ""),
        "collector": holding.get("collector", ""),
        "policy_number": holding.get("policy_number", ""),
        "insurer": product.insurer,
        "product": product.product,
        "conditions": product.conditions,
        "premium": int(holding.get("premium", 0)),
        "payment_label": holding.get("payment_label", ""),
        "share_rate": float(holding.get("share_rate", 100.0)),
        "recruiter_type": recruiter_type,
        "contract_date": holding.get("contract_date", ""),
        "status": holding.get("status", ""),
        "source_type": product.source_type,
        "first_year_rate": product.first_year_rate,
        "total_rate": product.total_rate,
        "sheet_name": product.sheet_name,
        "row_number": product.row_number,
    }


def _render_smart_product_picker(
    holding: dict,
    recommended: list[ProductRate],
    insurer_products: list[ProductRate],
    payout_rate: float,
    key_prefix: str,
) -> ProductRate | None:
    """추천은 최대 3개만, 직접 찾기는 동일 보험사의 전체 상품을 검색하도록 분리합니다."""
    recommended_groups = _product_groups(recommended)
    direct_groups = _product_groups(insurer_products)
    modes = ["추천 상품"] if recommended_groups else []
    modes.append("직접 찾기")
    mode = st.radio(
        "상품 연결 방법",
        modes,
        horizontal=True,
        key=f"{key_prefix}_mode",
        help="추천이 맞지 않으면 직접 찾기에서 같은 보험사의 전체 상품을 검색할 수 있습니다.",
    )
    groups = recommended_groups if mode == "추천 상품" else direct_groups
    if mode == "추천 상품":
        product_names = list(groups)[:3]
        label = f"추천 상품 · {len(product_names)}개"
        placeholder = "가장 적합한 상품을 선택해 주세요."
    else:
        product_names = _direct_product_names(holding, groups)
        label = f"{holding.get('insurer') or '보험사'} 상품 직접 찾기"
        placeholder = "상품명을 입력하거나 목록에서 선택해 주세요."

    selected_name = st.selectbox(
        label,
        product_names,
        index=0 if mode == "추천 상품" and len(product_names) == 1 else None,
        placeholder=placeholder,
        key=f"{key_prefix}_{'recommended' if mode == '추천 상품' else 'direct'}_product",
    ) if product_names else None
    all_condition_candidates = groups.get(selected_name, []) if selected_name else []
    payment_years = holding.get("payment_years")
    matched_payment_candidates = [
        product for product in all_condition_candidates
        if _payment_matches(product, payment_years)
    ] if payment_years is not None else all_condition_candidates
    condition_candidates = matched_payment_candidates
    condition_key = hashlib.sha1(str(selected_name).encode("utf-8")).hexdigest()[:10]
    if selected_name and payment_years is not None and not matched_payment_candidates:
        if len(all_condition_candidates) == 1:
            only_condition = all_condition_candidates[0]
            st.warning(
                "⚠️ 납입기간 불일치\n\n"
                f"원본 {payment_years}년납 · "
                f"수수료표 {_payment_condition_label(only_condition)}\n\n"
                "적용할 조건을 직접 확인해 주세요."
            )
            return st.selectbox(
                "납입기간 및 세부 조건 · 직접 확인",
                all_condition_candidates,
                index=None,
                placeholder="불일치 내용을 확인한 후 적용할 조건을 선택해 주세요.",
                format_func=lambda product: _condition_option_label(
                    product, payout_rate, all_condition_candidates
                ),
                key=f"{key_prefix}_condition_mismatch_{condition_key}",
            )
        st.warning(
            "⚠️ 납입기간 확인 필요\n\n"
            f"원본 {payment_years}년납과 일치하는 조건이 없습니다. "
            "다른 납기 조건을 확인해 주세요."
        )
        show_other_payment = st.checkbox(
            "다른 납기 조건 보기",
            value=False,
            key=f"{key_prefix}_other_payment_{condition_key}",
        )
        condition_candidates = all_condition_candidates if show_other_payment else []
        if not show_other_payment:
            st.selectbox(
                "납입기간 및 세부 조건",
                [f"{payment_years}년납 조건 없음 · 다른 납기 조건 보기를 선택해 주세요."],
                disabled=True,
                key=f"{key_prefix}_condition_payment_wait_{condition_key}",
            )
            return None
    if condition_candidates:
        label_prefix = f"원본 {payment_years}년납과 일치" if payment_years is not None and matched_payment_candidates else "납입기간 및 세부 조건"
        return st.selectbox(
            f"{label_prefix} · {len(condition_candidates)}개",
            condition_candidates,
            index=0 if len(condition_candidates) == 1 else None,
            placeholder="납입기간과 세부 조건을 선택해 주세요.",
            format_func=lambda product: _condition_option_label(
                product, payout_rate, condition_candidates
            ),
            key=f"{key_prefix}_condition_{condition_key}",
        )
    st.selectbox(
        "납입기간 및 세부 조건",
        ["상품을 선택하면 해당 상품의 조건만 표시됩니다."],
        disabled=True,
        key=f"{key_prefix}_condition_wait_{mode}",
    )
    return None


def _render_manual_entry(all_products: list[ProductRate]) -> None:
    with st.expander("계약 직접 추가", expanded=False):
        if not all_products:
            st.info("생보 또는 손보 수수료 예시표를 먼저 올려 주세요.")
            return
        source_options = [source for source in ("생보", "손보") if any(p.source_type == source for p in all_products)]
        source_type = st.radio("보험 구분", source_options, horizontal=True, key="manual_source_type")
        insurers = sorted({p.insurer for p in all_products if p.source_type == source_type})
        insurer = st.selectbox("보험회사", insurers, index=None, key="manual_insurer")
        products = [p for p in all_products if p.source_type == source_type and p.insurer == insurer]
        product_groups: dict[str, list[ProductRate]] = defaultdict(list)
        for product in products:
            display_name, _ = _product_display_parts(product.product)
            product_groups[display_name].append(product)
        product_names = sorted(product_groups)
        product_name = st.selectbox("상품", product_names, index=None, key="manual_product")
        candidates = _sort_condition_candidates(product_groups.get(product_name, []))
        if candidates:
            selected = st.selectbox(
                "납입기간 및 세부 조건", candidates, index=None,
                format_func=lambda p: _condition_option_label(p, candidates=candidates),
                key="manual_condition",
                placeholder="납입기간과 세부 조건을 선택해 주세요.",
            )
        else:
            selected = None
            st.selectbox(
                "납입기간 및 세부 조건",
                ["먼저 상품을 선택해 주세요."],
                disabled=True,
                key="manual_condition_disabled",
            )
        col1, col2 = st.columns(2)
        customer = col1.text_input("고객명", key="manual_customer")
        policy = col2.text_input("증권번호", key="manual_policy")
        premium = st.number_input("월보험료", min_value=0, step=1000, value=0, format="%d", key="manual_premium")
        if st.button("직접 입력 계약 추가", type="primary", use_container_width=True):
            if selected is None or premium <= 0:
                st.warning("보험회사·상품·세부 조건과 월보험료를 확인해 주세요.")
            else:
                holding = {
                    "customer": customer.strip(), "policy_number": policy.strip(), "premium": int(premium),
                    "share_rate": 100.0, "contract_date": "", "status": "직접 등록",
                }
                st.session_state["commission_contracts"].append(_contract_data(holding, selected))
                st.rerun()


def _render_contract_editor(all_products: list[ProductRate]) -> None:
    edit_index = st.session_state.get("commission_edit_index")
    contracts = st.session_state["commission_contracts"]
    if not isinstance(edit_index, int) or not (0 <= edit_index < len(contracts)):
        return

    contract = contracts[edit_index]
    st.markdown(
        '<div id="commission-edit-anchor" style="scroll-margin-top:5rem;"></div>',
        unsafe_allow_html=True,
    )
    components.html(
        """
        <script>
        setTimeout(function () {
            const target = window.parent.document.getElementById('commission-edit-anchor');
            if (target) {
                target.scrollIntoView({ behavior: 'smooth', block: 'start' });
            }
        }, 180);
        </script>
        """,
        height=0,
    )
    st.info(f"{edit_index + 1}번 계약을 수정하고 있습니다.")
    with st.container(border=True):
        customer_col, policy_col = st.columns(2)
        customer = customer_col.text_input(
            "고객명", value=contract.get("customer", ""), key=f"edit_customer_{edit_index}"
        )
        policy_number = policy_col.text_input(
            "증권번호", value=contract.get("policy_number", ""), key=f"edit_policy_{edit_index}"
        )
        premium_col, share_col = st.columns(2)
        premium = premium_col.number_input(
            "월보험료", min_value=0, step=1000, value=int(contract.get("premium", 0)),
            format="%d", key=f"edit_premium_{edit_index}",
        )
        share_rate = share_col.number_input(
            "쉐어율 (%)", min_value=0.0, max_value=100.0,
            value=float(contract.get("share_rate", 100.0)), step=1.0, format="%.0f",
            key=f"edit_share_{edit_index}",
        )

        source_options = [source for source in ("생보", "손보") if any(p.source_type == source for p in all_products)]
        current_source = contract.get("source_type")
        source_index = source_options.index(current_source) if current_source in source_options else 0
        selected_source = st.radio(
            "보험 구분", source_options, index=source_index, horizontal=True,
            key=f"edit_source_{edit_index}",
        )
        insurer_options = sorted({p.insurer for p in all_products if p.source_type == selected_source})
        current_insurer = contract.get("insurer")
        insurer_index = insurer_options.index(current_insurer) if current_insurer in insurer_options else None
        selected_insurer = st.selectbox(
            "보험회사", insurer_options, index=insurer_index,
            placeholder="보험회사를 선택해 주세요.", key=f"edit_insurer_{edit_index}",
        )

        insurer_products = [
            product for product in all_products
            if product.source_type == selected_source and product.insurer == selected_insurer
        ]
        product_groups: dict[str, list[ProductRate]] = defaultdict(list)
        for product in insurer_products:
            display_name, _ = _product_display_parts(product.product)
            product_groups[display_name].append(product)
        product_names = sorted(product_groups)
        current_product_name, _ = _product_display_parts(contract.get("product", ""))
        product_index = product_names.index(current_product_name) if current_product_name in product_names else None
        selected_product_name = st.selectbox(
            "상품", product_names, index=product_index,
            placeholder="상품을 선택해 주세요.", key=f"edit_product_{edit_index}_{selected_insurer}",
        ) if product_names else None

        matching = _sort_condition_candidates(product_groups.get(selected_product_name, []))
        current_position = next(
            (
                index for index, product in enumerate(matching)
                if product.sheet_name == contract.get("sheet_name") and product.row_number == contract.get("row_number")
            ),
            None,
        )
        if matching:
            selected_product = st.selectbox(
                "납입기간 및 세부 조건", matching,
                index=current_position,
                placeholder="납입기간과 세부 조건을 선택해 주세요.",
                format_func=lambda p: _condition_option_label(p, candidates=matching),
                key=f"edit_condition_{edit_index}_{hashlib.sha1(str(selected_product_name).encode()).hexdigest()[:8]}",
            )
        else:
            selected_product = None
            st.selectbox(
                "납입기간 및 세부 조건",
                ["먼저 상품을 선택해 주세요."],
                disabled=True,
                key=f"edit_condition_disabled_{edit_index}_{selected_insurer}",
            )

        recruiter_type = contract.get("recruiter_type", "")
        if share_rate < 100:
            recruiter_options = ["주모집", "공동모집"]
            recruiter_index = recruiter_options.index(recruiter_type) if recruiter_type in recruiter_options else None
            recruiter_type = st.selectbox(
                "모집 형태", recruiter_options, index=recruiter_index,
                placeholder="모집 형태를 선택해 주세요.", key=f"edit_recruiter_{edit_index}",
            ) or ""
        else:
            recruiter_type = ""

        save_col, cancel_col = st.columns([3, 1])
        if save_col.button("수정 완료", type="primary", use_container_width=True, key=f"save_edit_{edit_index}"):
            if premium <= 0:
                st.warning("월보험료를 확인해 주세요.")
            elif share_rate < 100 and not recruiter_type:
                st.warning("모집 형태를 선택해 주세요.")
            else:
                updated = dict(contract)
                updated.update({
                    "customer": customer.strip(), "policy_number": policy_number.strip(),
                    "premium": int(premium), "share_rate": float(share_rate),
                    "recruiter_type": recruiter_type,
                })
                if selected_product is not None:
                    updated.update({
                        "insurer": selected_product.insurer,
                        "product": selected_product.product,
                        "conditions": selected_product.conditions,
                        "source_type": selected_product.source_type,
                        "first_year_rate": selected_product.first_year_rate,
                        "total_rate": selected_product.total_rate,
                        "sheet_name": selected_product.sheet_name,
                        "row_number": selected_product.row_number,
                        "rate_recheck_required": False,
                    })
                contracts[edit_index] = updated
                st.session_state["commission_edit_index"] = None
                st.rerun()
        if cancel_col.button("취소", use_container_width=True, key=f"cancel_edit_{edit_index}"):
            st.session_state["commission_edit_index"] = None
            st.rerun()


def run() -> None:
    _initialize_state()

    page_header(
        "실적 관리",
        "수수료 계산기",
        "수수료 예시표와 보유계약 장기 파일을 연결해 계약별 예상 수당을 계산합니다.",
        "CC",
    )
    with st.container(horizontal=True):
        st.caption('① 수수료표 등록  →  ② 계약자료 확인  →  ③ 예상 수당 확인')
    with st.container(key="hw_surface_commission_calculator_0"):
        st.caption("01 · 수수료 예시표")
        section_intro("입력", "수수료 자료 불러오기", "생보·손보 수수료 예시표를 먼저 등록해 주세요.")

        with st.expander("① 수수료 예시표 불러오기", expanded=True):
            life_file = guarded_upload(
                "생보 수수료 예시표", type=["xlsx"], key="commission_life_file"
            )
            nonlife_file = guarded_upload(
                "손보 수수료 예시표", type=["xlsx"], key="commission_nonlife_file"
            )

        all_products: list[ProductRate] = []
        parse_warnings: list[str] = []
        reference_months: dict[str, str] = {}
        ratebook_hash = hashlib.sha256()
        for uploaded, source_type in ((life_file, "생보"), (nonlife_file, "손보")):
            if uploaded is None:
                continue
            try:
                uploaded_bytes = uploaded.getvalue()
                ratebook_hash.update(source_type.encode("utf-8"))
                ratebook_hash.update(uploaded_bytes)
                parsed, warnings = cached(
                    "parse:commission_calculator:rates:" + source_type,
                    fingerprint("commission-rates-v1", digest_bytes(uploaded_bytes), source_type, _month_from_filename(uploaded.name)),
                    lambda: parse_commission_workbook(uploaded_bytes, source_type),
                )
                all_products.extend(_to_product_rate(item) for item in parsed)
                parse_warnings.extend(warnings)
                reference_months[source_type] = _month_from_filename(uploaded.name)
            except Exception as exc:
                st.error("자료 처리에 실패했습니다. 파일 형식과 입력 내용을 확인한 뒤 다시 시도해 주세요. [PROCESS_FAILED]")
        if all_products:
            insurer_count = len({product.insurer for product in all_products})
            month_text = " · ".join(
                f"{source} {month.replace('-', '년 ')}월" if month else f"{source} 기준월 확인 필요"
                for source, month in reference_months.items()
            )
            st.success(f"{month_text} · 보험회사 {insurer_count}개 · 수수료 조건 {len(all_products):,}개")
        else:
            st.info("생보 또는 손보 수수료 예시표를 올리면 상품을 선택할 수 있습니다.")

        for warning in parse_warnings:
            st.warning(warning)

        current_ratebook_signature = ratebook_hash.hexdigest() if all_products else ""
        saved_ratebook_signature = st.session_state.get("commission_ratebook_signature", "")
        contracts = st.session_state["commission_contracts"]
        if current_ratebook_signature and not saved_ratebook_signature:
            st.session_state["commission_ratebook_signature"] = current_ratebook_signature
        elif (
            current_ratebook_signature and saved_ratebook_signature
            and current_ratebook_signature != saved_ratebook_signature and contracts
        ):
            st.warning(
                "수수료 예시표가 변경되었습니다. 기존 수당 계산 대상 계약의 요율을 "
                "새 수수료표 기준으로 재검증해야 합니다."
            )
            reconnect_col, clear_col = st.columns(2)
            if reconnect_col.button("새 수수료표로 다시 연결", type="primary", use_container_width=True):
                updated, unresolved = _reconnect_contract_rates(contracts, all_products)
                st.session_state["commission_ratebook_signature"] = current_ratebook_signature
                st.session_state["commission_edit_index"] = None
                st.toast(f"{updated}건 재연결 · {unresolved}건 직접 확인 필요")
                st.rerun()
            if clear_col.button("기존 계약 초기화", use_container_width=True):
                st.session_state["commission_contracts"] = []
                st.session_state["commission_ratebook_signature"] = current_ratebook_signature
                st.session_state["commission_edit_index"] = None
                st.rerun()
            st.stop()
        elif current_ratebook_signature and current_ratebook_signature != saved_ratebook_signature:
            st.session_state["commission_ratebook_signature"] = current_ratebook_signature

    with st.container(key="hw_surface_commission_calculator_1"):
        st.caption("02 · 계약과 지급 조건")
        section_intro("입력", "지급율 및 보유계약 불러오기", "공통 지급율을 확인하고 보유계약관리 장기 엑셀을 등록해 주세요.")
        payout_rate_percent = st.number_input(
            "공통 지급율 (%)",
            min_value=0.0,
            max_value=100.0,
            value=float(st.session_state["commission_payout_rate"]),
            step=0.1,
            format="%.12g",
            help="변경한 지급율은 현재 수당 계산 대상 계약에 일괄 적용됩니다.",
        )
        st.session_state["commission_payout_rate"] = payout_rate_percent
        payout_rate = payout_rate_percent / 100
        with st.expander('적용 지급률 확인', expanded=True):
            st.metric('현재 계산 지급률', f'{payout_rate_percent:g}%')
            st.caption('예시표의 회사별 수수료율에 이 지급률을 적용합니다. 상품·납입기간·예시표 기준월이 계약과 일치하는지 확인하세요.')
            from modules.shared.organization import active_profile
            current_policy = active_profile()
            st.caption(f"운영 기준: {current_policy['payout_rule_label']} · 버전 {current_policy['revision']} · 적용일 {current_policy['effective_from']}")
        holding_file = guarded_upload(
            "보유계약관리 장기 엑셀", type=["xlsx"], key="commission_holding_file",
            help="계약상태가 정상이고 수수료표 기준월과 같은 계약을 우선 분석합니다.",
        )

        # 선택할 때마다 전체 16,000여 조건을 다시 비교하지 않도록 보험사별로 미리 나눕니다.
        products_by_insurer: dict[tuple[str, str], list[ProductRate]] = defaultdict(list)
        for product in all_products:
            products_by_insurer[(product.source_type, product.insurer)].append(product)

        review_records: list[dict] = []
        if holding_file is not None and all_products:
            try:
                holdings = copy.deepcopy(cached(
                    "parse:commission_calculator:holdings",
                    fingerprint("holdings-v1", digest_bytes(holding_file.getvalue()), _month_from_filename(holding_file.name)),
                    lambda: parse_holding_workbook(holding_file.getvalue()),
                ))
            except Exception as exc:
                holdings = []
                st.error("자료 처리에 실패했습니다. 파일 형식과 입력 내용을 확인한 뒤 다시 시도해 주세요. [PROCESS_FAILED]")
            st.session_state["commission_import_collectors"] = list(dict.fromkeys(
                _clean_text(holding.get("collector", ""))
                for holding in holdings if _clean_text(holding.get("collector", ""))
            ))
            st.session_state["commission_import_contract_months"] = sorted({
                _clean_text(holding.get("contract_month", ""))
                for holding in holdings
                if re.fullmatch(r"20\d{2}-\d{2}", _clean_text(holding.get("contract_month", "")))
            })

            product_by_key = {product.key: product for product in all_products}
            product_rows = [product.__dict__ for product in all_products]
            link_decisions = _analyze_product_links(holdings, product_rows)
            registered_policies = {c.get("policy_number") for c in st.session_state["commission_contracts"] if c.get("policy_number")}
            automatic: list[tuple[dict, ProductRate]] = []
            needs_review: list[tuple[dict, list[ProductRate], str]] = []
            excluded: list[tuple[dict, str]] = []
            unmatched: list[tuple[dict, str]] = []
            already_registered = 0

            for holding in holdings:
                insurer_products = products_by_insurer.get(
                    (holding.get("source_type", ""), holding.get("insurer", "")), []
                )
                ref_month = reference_months.get(holding["source_type"], "")
                if holding.get("policy_number") and holding["policy_number"] in registered_policies:
                    already_registered += 1
                    continue
                if holding.get("status") != "정상":
                    excluded.append((holding, f"계약상태가 {holding.get('status') or '확인 필요'}이므로 기본 제외"))
                    continue
                if ref_month and holding.get("contract_month") and holding["contract_month"] != ref_month:
                    excluded.append((holding, f"계약월 {holding['contract_month']} / 수수료표 기준월 {ref_month}"))
                    continue
                decision = link_decisions.get(holding["row_key"], {})
                candidates = [
                    product_by_key[key] for key in decision.get("candidate_keys", [])
                    if key in product_by_key
                ]
                if not candidates:
                    unmatched.append((holding, "수수료표에서 일치하는 상품을 찾지 못함"))
                    continue
                auto = product_by_key.get(decision.get("auto_key", ""))
                if auto is not None and holding.get("share_rate", 100.0) >= 100:
                    automatic.append((holding, auto))
                else:
                    reason_parts = []
                    if auto is None:
                        reason_parts.append("세부 조건 확인")
                    if holding.get("share_rate", 100.0) < 100:
                        reason_parts.append("모집 형태 확인")
                    review_candidates = [
                        product_by_key[key] for key in decision.get("review_keys", [])
                        if key in product_by_key
                    ]
                    needs_review.append((holding, review_candidates, " · ".join(reason_parts)))

            section_intro("연결 결과", "자동 연결 및 확인 필요 계약", "자동 연결 결과를 검토하고 필요한 계약만 조건을 다시 확인해 주세요.")
            metric_cols = st.columns(4)
            metric_cols[0].metric("전체", f"{len(holdings)}건")
            metric_cols[1].metric("자동 연결", f"{len(automatic)}건")
            metric_cols[2].metric("확인 필요", f"{len(needs_review)}건")
            metric_cols[3].metric("미연결·제외", f"{len(unmatched) + len(excluded)}건")
            if already_registered:
                st.caption(f"이미 등록된 증권번호 {already_registered}건은 중복 분석에서 제외했습니다.")

            pending: list[dict] = []
            with st.expander(f"자동 연결 완료 {len(automatic)}건", expanded=True):
                if not automatic:
                    st.caption("자동 연결된 계약이 없습니다.")
                for holding, product in automatic:
                    col1, col2 = st.columns([0.08, 0.92])
                    selected = col1.checkbox("선택", value=True, key=f"auto_{holding['row_key']}", label_visibility="collapsed")
                    selected_product = product
                    with col2:
                        customer_name = _markdown_text(holding.get("customer") or "고객명 없음")
                        st.markdown(f"**{customer_name} · {product.insurer}**")
                        st.caption(_holding_caption(holding))
                        st.write(f"{product.product} · {product.conditions or '기본 조건'}")
                        reason = "상품명 일치"
                        if holding.get("payment_label"):
                            reason += f" · {holding['payment_label']} 조건 일치"
                        _, _, tag_reasons = _tag_match_summary(holding, product)
                        if tag_reasons:
                            reason += " · " + " · ".join(dict.fromkeys(tag_reasons))
                        st.caption(f"자동 연결 근거: {reason}")
                        verify_auto = st.checkbox(
                            "상품·납기 다시 확인",
                            value=False,
                            key=f"auto_verify_{holding['row_key']}",
                        )
                        if verify_auto:
                            insurer_products = products_by_insurer.get(
                                (holding.get("source_type", ""), holding.get("insurer", "")), []
                            )
                            decision = link_decisions.get(holding["row_key"], {})
                            recommended = [
                                product_by_key[key] for key in decision.get("review_keys", [])
                                if key in product_by_key
                            ]
                            selected_product = _render_smart_product_picker(
                                holding,
                                recommended or [product],
                                insurer_products,
                                payout_rate,
                                key_prefix=f"auto_change_{holding['row_key']}",
                            )
                    if selected and selected_product is not None:
                        pending.append(_contract_data(holding, selected_product))
                    elif selected and verify_auto:
                        st.caption("변경할 상품과 원본 납기에 맞는 조건을 선택해 주세요.")

            with st.expander(f"확인 필요 {len(needs_review)}건", expanded=bool(needs_review)):
                if not needs_review:
                    st.caption("확인이 필요한 계약이 없습니다.")
                for holding, candidates, reason in needs_review:
                    with st.container(border=True):
                        customer_name = _markdown_text(holding.get("customer") or "고객명 없음")
                        st.markdown(f"**{customer_name} · {holding['insurer']}**")
                        st.caption(f"{_holding_caption(holding)} · {reason}")
                        st.write(f"보유계약 상품: {holding['product_raw']}")
                        insurer_products = products_by_insurer.get(
                            (holding.get("source_type", ""), holding.get("insurer", "")), []
                        )
                        selected_product = _render_smart_product_picker(
                            holding, candidates, insurer_products, payout_rate,
                            key_prefix=f"review_{holding['row_key']}",
                        )
                    recruiter_type = ""
                    if holding.get("share_rate", 100.0) < 100:
                        recruiter_type = st.selectbox(
                            f"모집 형태 · 엑셀 쉐어율 {holding['share_rate']:g}%",
                            ["주모집", "공동모집"], index=None, key=f"recruiter_{holding['row_key']}",
                            placeholder="모집 형태를 선택해 주세요.",
                        ) or ""
                    include = st.checkbox("이 계약 등록", value=True, key=f"review_include_{holding['row_key']}")
                    ready = include and selected_product is not None and (
                        holding.get("share_rate", 100.0) >= 100 or recruiter_type
                    )
                    if ready:
                        pending.append(_contract_data(holding, selected_product, recruiter_type))
                        st.success("등록 준비 완료")
                    elif not include:
                        review_records.append({**holding, "product": holding["product_raw"], "reason": "사용자가 등록 대상에서 제외"})
                    else:
                        st.caption("상품·납입기간·모집 형태 중 필요한 항목을 선택해 주세요.")
                    st.write("")

            if excluded:
                with st.expander(f"기준월·계약상태·중복으로 제외 {len(excluded)}건", expanded=False):
                    st.info("기본적으로 제외됩니다. 필요한 경우에만 계약을 펼쳐 포함해 주세요.")
                    for holding, reason in excluded:
                        customer_name = _markdown_text(holding.get("customer") or "고객명 없음")
                        st.markdown(f"**{customer_name} · {holding['insurer']}**")
                        st.caption(f"{_holding_caption(holding)} · {reason}")
                        include = st.checkbox("이번 계산에 포함", value=False, key=f"excluded_include_{holding['row_key']}")
                        if include:
                            insurer_products = products_by_insurer.get(
                                (holding.get("source_type", ""), holding.get("insurer", "")), []
                            )
                            candidates = _candidate_products(holding, insurer_products)
                            selected_product = st.selectbox(
                                "적용할 상품 및 조건", candidates, index=None, key=f"excluded_product_{holding['row_key']}",
                                format_func=lambda p: (
                                    f"{p.product} · "
                                    f"{_condition_option_label(p, payout_rate, candidates)}"
                                ),
                            ) if candidates else None
                            confirmed = st.checkbox(
                                "제외 사유를 확인했으며 이번 계산에 포함합니다.",
                                key=f"excluded_confirm_{holding['row_key']}",
                            )
                            if selected_product is not None and confirmed:
                                pending.append(_contract_data(holding, selected_product))
                            else:
                                review_records.append({**holding, "product": holding["product_raw"], "reason": reason})
                        else:
                            review_records.append({**holding, "product": holding["product_raw"], "reason": reason})
                        st.divider()

            if unmatched:
                with st.expander(f"연결되지 않은 계약 {len(unmatched)}건", expanded=bool(unmatched)):
                    for holding, reason in unmatched:
                        customer_name = _markdown_text(holding.get("customer") or "고객명 없음")
                        st.markdown(f"**{customer_name} · {holding['insurer']}**")
                        st.caption(_holding_caption(holding))
                        st.write(f"{holding['product_raw']} · {reason}")
                        with st.container(border=True):
                            insurer_products = products_by_insurer.get(
                                (holding.get("source_type", ""), holding.get("insurer", "")), []
                            )
                            if not insurer_products:
                                source_products = [
                                    product for product in all_products
                                    if product.source_type == holding.get("source_type")
                                ]
                                insurer_options = sorted({product.insurer for product in source_products})
                                direct_insurer = st.selectbox(
                                    "보험회사를 찾지 못했습니다 · 직접 선택",
                                    insurer_options,
                                    index=None,
                                    placeholder="보험회사를 선택해 주세요.",
                                    key=f"unmatched_insurer_{holding['row_key']}",
                                )
                                insurer_products = [
                                    product for product in source_products
                                    if product.insurer == direct_insurer
                                ]
                            direct_product = _render_smart_product_picker(
                                holding, [], insurer_products, payout_rate,
                                key_prefix=f"unmatched_{holding['row_key']}",
                            )
                            if direct_product is not None:
                                pending.append(_contract_data(holding, direct_product))
                                st.success("직접 연결 준비 완료")
                            else:
                                review_records.append({
                                    **holding, "product": holding["product_raw"], "reason": reason
                                })
                        st.divider()

            if pending:
                if st.button(f"선택한 계약 {len(pending)}건 등록", type="primary", use_container_width=True):
                    existing = {c.get("policy_number") for c in st.session_state["commission_contracts"] if c.get("policy_number")}
                    added = 0
                    for contract in pending:
                        if contract.get("policy_number") and contract["policy_number"] in existing:
                            continue
                        st.session_state["commission_contracts"].append(contract)
                        if contract.get("policy_number"):
                            existing.add(contract["policy_number"])
                        added += 1
                    st.toast(f"계약 {added}건을 등록했습니다.")
                    st.rerun()
            elif holdings:
                st.info("현재 등록할 수 있는 계약이 없습니다. 확인 필요 계약의 조건을 선택해 주세요.")

        section_intro("직접 입력", "계약 직접 추가", "파일에 없는 계약은 보험회사와 상품 조건을 직접 선택해 추가할 수 있습니다.")
        _render_manual_entry(all_products)

        contracts = st.session_state["commission_contracts"]
    with st.container(key="hw_surface_commission_calculator_2"):
        st.caption("03 · 예상 수당 · 계약별 결과")
        section_intro("계산 결과", "수당 계산 대상 계약", "등록된 계약과 예상 익월수당·총수당을 확인해 주세요.")
        if not contracts:
            st.info("추가된 계약이 없습니다.")
            return

        _render_contract_editor(all_products)

        calculation_contracts = [contract for contract in contracts if not contract.get("rate_recheck_required")]
        recheck_count = len(contracts) - len(calculation_contracts)
        if recheck_count:
            st.warning(
                f"새 수수료표에서 요율을 확정하지 못한 계약 {recheck_count}건은 합계와 다운로드에서 제외했습니다. "
                "해당 계약의 수정 버튼을 눌러 상품과 세부 조건을 다시 선택해 주세요."
            )

        total_premium = sum(contract["premium"] for contract in calculation_contracts)
        total_first = sum(
            contract["premium"] * contract["first_year_rate"] * payout_rate
            for contract in calculation_contracts
        )
        total_commission = sum(
            contract["premium"] * contract["total_rate"] * payout_rate
            for contract in calculation_contracts
        )
        metric_cols = st.columns(3)
        metric_cols[0].metric("월보험료 합계", _format_won(total_premium))
        metric_cols[1].metric("예상 익월수당", _format_won(total_first))
        metric_cols[2].metric("예상 총수당", _format_won(total_commission))

        header_columns = st.columns([3.6, 1, 1.15, 1.15, 1.25, 1.25, 1.05])
        for column, label in zip(
            header_columns, ("계약 정보", "월보험료", "익월 수수료율", "총 수수료율", "예상 익월수당", "예상 총수당", "관리"),
        ):
            column.caption(label)

        for index, contract in enumerate(contracts):
            first_rate = contract["first_year_rate"] * payout_rate
            total_rate = contract["total_rate"] * payout_rate
            expected_first = contract["premium"] * first_rate
            expected_total = contract["premium"] * total_rate
            product_detail = _compact_product_display(contract, contracts)

            row_columns = st.columns([3.6, 1, 1.15, 1.15, 1.25, 1.25, 1.05])
            with row_columns[0]:
                customer_name = _markdown_text(contract.get("customer") or "고객명 없음")
                st.markdown(f"**{index + 1}. {customer_name}** · {contract['insurer']}")
                policy = contract.get("policy_number") or "증권번호 없음"
                recruiting = ""
                if contract.get("share_rate", 100) < 100:
                    recruiting = f" · {contract['share_rate']:g}% · {contract.get('recruiter_type') or '모집 형태 확인'}"
                st.caption(f"증권번호 {policy}{recruiting}")
                st.markdown(_markdown_text(product_detail).replace("\n", "  \n"))
                if contract.get("rate_recheck_required"):
                    st.warning("요율 재확인 필요")
            row_columns[1].write(_format_won(contract["premium"]))
            row_columns[2].write(_format_rate(first_rate))
            row_columns[3].write(_format_rate(total_rate))
            row_columns[4].write(_format_won(expected_first))
            row_columns[5].write(_format_won(expected_total))
            with row_columns[6]:
                edit_col, delete_col = st.columns(2)
                if edit_col.button("수정", key=f"edit_commission_{index}", help="이 계약 수정"):
                    st.session_state["commission_edit_index"] = index
                    st.rerun()
                if delete_col.button("✕", key=f"delete_commission_{index}", help="이 계약 삭제"):
                    contracts.pop(index)
                    current_edit = st.session_state.get("commission_edit_index")
                    if current_edit == index:
                        st.session_state["commission_edit_index"] = None
                    elif isinstance(current_edit, int) and current_edit > index:
                        st.session_state["commission_edit_index"] = current_edit - 1
                    st.rerun()

            if index < len(contracts) - 1:
                st.markdown(
                    '<hr style="margin:.25rem 0 .45rem;border:0;border-top:1px solid rgba(128,128,128,.18);">',
                    unsafe_allow_html=True,
                )

        section_intro("다운로드", "계산 결과 내려받기", "확인된 계약과 수수료 계산 결과를 엑셀로 저장합니다.")
        clear_col, download_col = st.columns([1, 2])
        with clear_col:
            if st.button("전체 계약 지우기", use_container_width=True):
                st.session_state["commission_contracts"] = []
                st.session_state["commission_edit_index"] = None
                st.rerun()
        with download_col:
            months = sorted({month for month in reference_months.values() if month})
            reference_month = ", ".join(months)
            excel_bytes = _make_excel(
                calculation_contracts,
                payout_rate,
                reference_month,
                review_records,
                st.session_state.get("commission_import_collectors", []),
            )
            st.download_button(
                "엑셀 다운로드",
                data=excel_bytes,
                file_name=_commission_download_filename(
                    calculation_contracts,
                    st.session_state.get("commission_import_collectors", []),
                    st.session_state.get("commission_import_contract_months", []),
                ),
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary",
                use_container_width=True,
            )
    page_footer("수수료 계산기", APP_VERSION)
