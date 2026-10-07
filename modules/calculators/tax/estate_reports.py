"""Single calculation snapshot for customer PDF and detailed workbook."""
from io import BytesIO
from html import escape
from decimal import Decimal
from modules.shared.runtime_cache import session_export


def money(v):
    return f'{Decimal(str(v)):,.0f}원'


def headline(result, funding):
    due=result.metrics['상속세 추정액']
    if funding is None:
        return f'예상 상속세는 {money(due)}입니다. 실제 사용할 수 있는 납부재원을 입력하면 부족한 금액을 함께 확인할 수 있습니다.'
    total=sum(funding)
    if due == 0:return '현재 입력 조건에서 예상 상속세는 없습니다.'
    if total < due:return f'예상 상속세 {money(due)}에 준비된 납부재원 {money(total)}을 반영하면, {money(due-total)}의 추가 재원 마련이 필요합니다.'
    return f'입력한 납부재원으로 예상 상속세를 충당하면 {money(total-due)}의 여유가 예상됩니다.'


def metrics(result,funding):
    due=result.metrics['상속세 추정액']
    if funding is None:return [('예상 상속세',money(due)),('준비된 납부재원','미입력'),('납부재원 부족액','확인 필요')]
    gap=sum(funding)-due
    return [('예상 상속세',money(due)),('준비된 납부재원',money(sum(funding))),('납부재원 부족액' if gap<0 else '납부 후 예상 여유액',money(abs(gap)))]


def summary_rows(result):
    keys=['상속세 과세가액','실제 적용 공제','과세표준','산출세액','세대생략 할증','증여·기타 세액공제','신고세액공제','상속세 추정액']
    return [(row['항목'],money(row['금액'])) for key in keys for row in result.rows if row['항목']==key and (row['금액'] or key in keys[:4]+keys[-1:])]


@session_export("estate-pdf-v2", name_arg=False)
def pdf_report(result, conditions, funding, stamp, alias='', note='', alignment='왼쪽'):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle
    from modules.shared.pdf_brand import draw_brand
    from modules.shared.report_fonts import korean_pdf_font
    font = korean_pdf_font()
    navy=colors.HexColor('#112B49'); pale=colors.HexColor('#EEF4F8')
    body=ParagraphStyle('estate',fontName=font,fontSize=9,leading=13,wordWrap='CJK',textColor=navy)
    def p(v,style=body):return Paragraph(escape(str(v)).replace('\n','<br/>'),style)
    buf=BytesIO();doc=SimpleDocTemplate(buf,pagesize=A4,leftMargin=34,rightMargin=34,topMargin=56,bottomMargin=47,title='상속세·납부재원 요약',author='화랑 WORKSPACE')
    width=A4[0]-68
    story=[p('상속세·납부재원 요약',ParagraphStyle('title',parent=body,fontSize=20,leading=26)),Spacer(1,6),p('작성일 '+stamp+('  |  '+alias if alias else '')),Spacer(1,12)]
    def table(rows,card=False):
        if card:
            center=ParagraphStyle('center',parent=body,alignment=1)
            data=[[p(a,center) for a,b in rows],[p(b,ParagraphStyle('big',parent=center,fontSize=13,leading=19)) for a,b in rows]]
            t=Table(data,colWidths=[width/3]*3)
        else:
            right=ParagraphStyle('right',parent=body,alignment=2)
            t=Table([[p(a),p(b,right)] for a,b in rows],colWidths=[width*.55,width*.45],splitByRow=1,splitInRow=1)
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),pale),('VALIGN',(0,0),(-1,-1),'TOP'),('LINEBELOW',(0,0),(-1,-1),.4,colors.HexColor('#DCE5EF')),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6),('LEFTPADDING',(0,0),(-1,-1),9),('RIGHTPADDING',(0,0),(-1,-1),9)]));story.append(t)
    def heading(s):story.extend([Spacer(1,10),p(s,ParagraphStyle('head',parent=body,fontSize=11,leading=16,keepWithNext=True)),Spacer(1,4)])
    table(metrics(result,funding),True);story.extend([Spacer(1,8),p(headline(result,funding))])
    heading('계산 조건');story.append(p(' · '.join(f'{a}: {b}' for a,b in conditions)))
    heading('세액 계산 요약');table(summary_rows(result))
    if funding is not None:
        heading('납부재원 구성');table(list(zip(['현금·예금','납부에 사용할 사망보험금','기타 자금','총 납부재원'],[money(x) for x in (*funding,sum(funding))])))
    if note.strip():
        heading('고객님께 드리는 설명')
        style=ParagraphStyle('note',parent=body,alignment={'왼쪽':0,'가운데':1,'오른쪽':2}[alignment])
        for line in note.split('\n'):story.append(p(line,style) if line else Spacer(1,6))
    heading('계산 기준')
    from modules.calculators.tax.personal_tax_models import AS_OF
    story.append(p(f'계산 기준일 {AS_OF} · 2026년 거주자 상속. 입력 조건과 확인된 공제 요건에 따른 예상치이며 실제 신고세액과 다를 수 있습니다. 납부재원은 실제 사용 가능한 금액을 기준으로 합니다.'))
    doc.build(story,onFirstPage=lambda c,d:draw_brand(c,d.page),onLaterPages=lambda c,d:draw_brand(c,d.page))
    return buf.getvalue()


@session_export("estate-excel-v1", name_arg=False)
def excel_report(result,inputs,conditions,funding,stamp,note=''):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    wb=Workbook();wb.remove(wb.active)
    rows=[(r['항목'],r['금액']) for r in result.rows]
    deductions=[r for r in rows if '공제' in r[0] or '한도' in r[0]]
    sheets={'요약':[('작성일',stamp),*conditions,*metrics(result,funding),('결과 설명',headline(result,funding)),('고객에게 전할 설명',note)],'공제 상세':deductions,'세액 계산 상세':rows,'입력 내역':inputs}
    if funding is not None:sheets['요약']+=list(zip(['현금·예금 (원)','사망보험금 (원)','기타 자금 (원)'],funding))
    for title,data in sheets.items():
        ws=wb.create_sheet(title);ws.append(['화랑 WORKSPACE',title]);ws.append(['항목','금액(원) / 내용'])
        for a,b in data:ws.append([a,float(b) if isinstance(b,Decimal) else b])
        ws.column_dimensions['A'].width=55;ws.column_dimensions['B'].width=75
        for row in ws:
            for cell in row:
                cell.font=Font(name='맑은 고딕',size=11,color='112B49');cell.alignment=Alignment(vertical='top',wrap_text=True)
                if cell.row<=2:cell.fill=PatternFill('solid',fgColor='112B49');cell.font=Font(name='맑은 고딕',bold=True,color='FFFFFF')
                elif isinstance(cell.value,(int,float)):cell.number_format='#,##0';cell.alignment=Alignment(horizontal='right')
            ws.row_dimensions[row[0].row].height=32 if row[0].row<=2 else 44
        ws.freeze_panes='B3';ws.print_title_rows='1:2';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.orientation='landscape';ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.oddHeader.right.text='화랑 WORKSPACE';ws.oddFooter.right.text='&P'
    b=BytesIO();wb.save(b);return b.getvalue()
