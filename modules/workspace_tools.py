"""Session-only UI helpers and in-memory, literal-text exports."""
from __future__ import annotations

import html
import io
import re
from pathlib import Path

import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from .session_store import commit_input, reset_page

CHECKED = "2026-09-21"
SOURCES = {
    "보험나이": "https://www.cardif.co.kr/customer-center/보험나이-만나이.do",
    "생명보험협회": "https://www.klia.or.kr/",
    "손해보험협회": "https://www.knia.or.kr/",
    "국민건강보험": "https://www.nhis.or.kr/nhis/index.do",
    "AIA 청구서식": "https://www.aia.co.kr/ko/customer-support/customer-guide/forms/claims.html",
    "iM라이프 청구서식": "https://www.imlifeins.co.kr/BB/BB_D030.do",
    "설명 참고": "https://carinfo.knia.or.kr/lmxsrv/law/lawFullContent.do?SEQ=4&SEQ_HISTORY=9",
}

_PREFIX_PAGES = {
    "a_": "quick_calculators", "b_": "consultation_helper",
    "c_": "comparison_builder", "e_": "education_center", "f_": "insurer_portal",
    "d_": "customer_materials",
}


def text(value, limit=2000):
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str("" if value is None else value))[:limit]


def field(kind, label, key, value=None, **kwargs):
    """Keep values across page changes without global/disk caches."""
    ui = "_ws_" + key
    if key not in st.session_state:
        st.session_state[key] = value
    if ui not in st.session_state:
        st.session_state[ui] = st.session_state[key]
    def remember():
        st.session_state[key] = st.session_state[ui]
        if key.startswith('a_'):
            st.session_state['a_review_token'] = None
        if key.startswith('d_'):
            st.session_state['d_review'] = None
        page = next((page for prefix, page in _PREFIX_PAGES.items() if key.startswith(prefix)), None)
        if page:
            commit_input(page, key, st.session_state[ui])
    return getattr(st, kind)(label, key=ui, on_change=remember, **kwargs)


def clear_namespace(prefix):
    page = _PREFIX_PAGES.get(prefix)
    if page:
        reset_page(page)
        return
    for key in list(st.session_state):
        if key.startswith(prefix) or key.startswith("_ws_" + prefix):
            del st.session_state[key]


def session_notice(prefix):
    with st.expander("입력 초기화", expanded=False):
        st.button("이 도구 입력 초기화", key="clear_" + prefix, on_click=clear_namespace, args=(prefix,))


def source_notes(names):
    with st.expander("기준과 출처"):
        st.caption(f"콘텐츠 편집 기준일 {CHECKED} · 일반 참고자료. 링크별 접속·최신 개정 확인일을 의미하지 않습니다. 개별 약관·회사 승인 자료를 우선 확인하세요.")
        for name in names:
            st.link_button(name, SOURCES[name], use_container_width=True)


def workbook_bytes(title, headers, rows, notes=""):
    wb = Workbook()
    ws = wb.active
    ws.title = "상담 비교"
    ws.append([text(title)])
    ws.append(["HWARANG · 상담 참고용 · 입력값 기반"])
    ws.append([text(notes, 16000)])
    ws.append(list(headers))
    for row in rows:
        ws.append([text(v, 1500) if not isinstance(v, (int, float)) else v for v in row])
    for row in ws:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # '=...', '+...', '@...' remain literal text.
            cell.font = Font(name="맑은 고딕", size=11, color="17233C")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in ws[4]:
        cell.fill = PatternFill("solid", fgColor="17233C")
        cell.font = Font(name="맑은 고딕", size=11, color="FFFFFF", bold=True)
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[ws.cell(4, i).column_letter].width = 26
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{ws.cell(max(4, ws.max_row), len(headers)).coordinate}"
    ws.oddFooter.center.text = "화랑 WORKSPACE | 상담 참고용 | &P / &N"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def pdf_bytes(title, headers, rows, notes=""):
    font = "WorkspacePretendard"
    if font not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font, str(Path(__file__).resolve().parents[1] / "assets/fonts/PretendardVariable.ttf")))
    output = io.BytesIO()
    page = landscape(A4)
    doc = SimpleDocTemplate(output, pagesize=page, rightMargin=30, leftMargin=30, topMargin=36, bottomMargin=40)
    style = ParagraphStyle("body", fontName=font, fontSize=10, leading=15, wordWrap="CJK", textColor=colors.HexColor("#17233C"))
    heading = ParagraphStyle("heading", parent=style, fontSize=20, leading=27, spaceAfter=12)
    cell = lambda value: Paragraph(html.escape(text(value, 16000)).replace("\n", "<br/>"), style)
    table = Table([[cell(v) for v in headers]] + [[cell(v) for v in row] for row in rows], colWidths=[(doc.width-12)/len(headers)]*len(headers), repeatRows=1, hAlign="LEFT", splitInRow=1)
    table.setStyle(TableStyle([("BACKGROUND", (0,0),(-1,0), colors.HexColor("#ECE8DF")), ("GRID",(0,0),(-1,-1),.5, colors.HexColor("#D4D6DB")), ("VALIGN",(0,0),(-1,-1),"TOP"), ("TOPPADDING",(0,0),(-1,-1),8), ("BOTTOMPADDING",(0,0),(-1,-1),8)]))
    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont(font, 9)
        canvas.setFillColor(colors.HexColor("#586277"))
        canvas.drawString(30, 20, f"H  |  화랑 WORKSPACE · 상담 참고용 · {document.page}")
        canvas.setFillColor(colors.Color(.72,.65,.48,alpha=.11))
        canvas.setFont(font, 45)
        canvas.translate(page[0]/2, page[1]/2)
        canvas.rotate(25)
        canvas.drawCentredString(0,0,"HWARANG · 상담 참고용")
        canvas.restoreState()
    note_paragraphs=[cell(part) for part in text(notes,16000).split("\n") if part.strip()]
    story=[Paragraph(html.escape(text(title,100)), heading), *note_paragraphs, Spacer(1,12),table]
    doc.build(story,onFirstPage=footer,onLaterPages=footer)
    return output.getvalue()


def downloads(prefix, title, headers, rows, notes=""):
    a,b=st.columns(2)
    with a:
        st.download_button("Excel 내려받기",workbook_bytes(title,headers,rows,notes),f"{prefix}_reference.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",key=prefix+"_xlsx",use_container_width=True)
    with b:
        st.download_button("PDF 내려받기",pdf_bytes(title,headers,rows,notes),f"{prefix}_reference.pdf","application/pdf",key=prefix+"_pdf",use_container_width=True)
