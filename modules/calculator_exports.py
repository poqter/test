"""Identical formatted values in UI, TXT, Excel and PDF; memory-only."""
import io
import math
import unicodedata
from openpyxl import load_workbook
from .calculator_core import won
from .consultation_documents import document_pdf
from .workspace_tools import workbook_bytes

NOTE = '사용자 입력과 가정에 따른 상담 참고자료입니다. 수익·지급을 보장하거나 권장 가입금액을 판정하지 않습니다.'


def formatted_results(results):
    return [(label, value if isinstance(value, str) else f'{won(value):,}원') for label, value in results.items()]


def sections(result):
    return [('입력 조건', result['inputs']), ('계산 결과', formatted_results(result['values'])),
            ('계산 근거', [('산식', result['formula'])]), ('가정과 미반영 조건', [('확인사항', result['assumptions'])])]


def text_report(result):
    return result['title']+'\n작성일 '+result['prepared_on']+'\n\n'+'\n\n'.join(
        title+'\n'+'\n'.join(f'{k}: {v}' for k,v in rows) for title,rows in sections(result))+'\n\n'+NOTE


def export_bytes(result, kind):
    if kind == 'txt': return text_report(result).encode('utf-8-sig')
    if kind == 'xlsx':
        rows = [(section, label, value) for section, items in sections(result) for label,value in items]
        wb = load_workbook(io.BytesIO(workbook_bytes(result['title'], ['구분','항목','내용'], rows, '작성일 '+result['prepared_on']+' · '+NOTE)))
        ws = wb.active
        ws.column_dimensions['A'].width = 24
        ws.column_dimensions['B'].width = 38
        ws.column_dimensions['C'].width = 75
        for row in (1,2,3):
            ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=3)
        ws.row_dimensions[1].height = 30
        ws.row_dimensions[3].height = 44
        for row in ws.iter_rows(min_row=5):
            lines = max(math.ceil(sum(2 if unicodedata.east_asian_width(c) in 'WF' else 1 for c in str(cell.value or '')) / (width-3)) for cell,width in zip(row,(24,38,75)))
            ws.row_dimensions[row[0].row].height = max(30,lines*17+10)
        ws.print_options.horizontalCentered = True
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = '1:4'
        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()
    if kind == 'pdf':
        body = '\n\n'.join(section+'\n'+'\n'.join(f'{label}: {value}' for label,value in items) for section,items in sections(result))
        return document_pdf(result['title'], body, prepared_on=result['prepared_on'], note=NOTE)
    raise ValueError('Unknown export format')
