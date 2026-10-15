"""Anonymous workbook in the exact accepted analyzer layout; never a real contract."""
from io import BytesIO
from openpyxl import Workbook
from functools import lru_cache

@lru_cache(maxsize=1)
def sample_workbook():
    book=Workbook();contracts=book.active;contracts.title='계약사항';coverage=book.create_sheet('상품별보장내용')
    contracts['B2']='가상고객을 위한 학습 예시';contracts['D2']='40세'
    for col,company,product,monthly in [(6,'가상 생명','학습용 보장 A',100000),(7,'가상 손해','학습용 보장 B',50000)]:
        for row,value in [(2,company),(3,product),(4,'90세'),(5,'12/240회'),(6,'월납'),(7,monthly)]:coverage.cell(row,col,value)
        r=9+col-6
        for c,value in [(10,monthly*240),(11,monthly*12),(12,monthly*228)]:contracts.cell(r,c,value)
    for row,label,amounts in [(9,'일반암진단비',[30000000,10000000]),(10,'질병입원일당',[30000,0]),(11,'학습용 추가 확인 담보',[0,1000000])]:
        coverage.cell(row,2,label)
        for col,value in enumerate(amounts,6):coverage.cell(row,col,value)
    info=book.create_sheet('읽어주세요');info['A1']='실제 보험사·상품·계약이 아닌 익명 학습 예시입니다.';info['A2']='필수 시트 계약사항 / 상품별보장내용의 행·열 배치를 유지하세요.'
    out=BytesIO();book.save(out);book.close();return out.getvalue()
