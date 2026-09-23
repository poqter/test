import io
import unittest
import zipfile
import xml.etree.ElementTree as ET
from unittest.mock import patch
from openpyxl import Workbook,load_workbook
from modules.analyzer import build_analysis_file
from modules.validators import validate_upload,ValidationError
from fixture_factory import synthetic_xlsx


def modified_archive(transform):
    output=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(synthetic_xlsx())) as source,zipfile.ZipFile(output,'w') as target:
        for name in source.namelist():
            data=transform(name,source.read(name))
            if data is not None:target.writestr(name,data)
    return output.getvalue()


class ReleaseRecheckTests(unittest.TestCase):
    def test_all_contract_text_and_intended_formulas(self):
        for value in ['=2+2','+SUM(1,2)','@SUM(1,2)','-2+2','일반상품']:
            with self.subTest(value=value):
                w=Workbook();c=w.active;c.title='계약사항';c['B2']='가상고객';c['D2']='40세'
                c['J9'],c['K9'],c['L9']=24000000,12000000,12000000
                c=w.create_sheet('상품별보장내용')
                for row in range(2,7):c.cell(row,6,value).data_type='s'
                c['F7']=100000;c['B9']='일반암';c['F9']=5000
                output=io.BytesIO();w.save(output)
                binary,_,_=build_analysis_file(output.getvalue(),selected_labels=['일반암'])
                result=load_workbook(io.BytesIO(binary));s=result['보장 분석']
                for row in range(2,7):
                    self.assertEqual(s.cell(row,4).value,value)
                    self.assertEqual(s.cell(row,4).data_type,'s')
                self.assertEqual(s['A7'].value,'=SUM(D7:D7)')
                self.assertEqual(s['A7'].data_type,'f')
                self.assertEqual(s['D7'].value,100000)
                self.assertEqual(s['D12'].value,5000)
                self.assertTrue(s.print_area)
                self.assertEqual(result.sheetnames,['보장 분석','보장 제안서'])

    def test_valid_workbook(self):
        validate_upload('valid.xlsx',synthetic_xlsx())

    def test_fake_xml_roots(self):
        data=modified_archive(lambda name,data:b'<dummy/>' if name=='xl/workbook.xml' else data)
        with self.assertRaises(ValidationError):validate_upload('fake.xlsx',data)

    def test_missing_sheet_target(self):
        data=modified_archive(lambda name,data:None if name=='xl/worksheets/sheet1.xml' else data)
        with self.assertRaises(ValidationError):validate_upload('missing.xlsx',data)

    def test_external_sheet_relation(self):
        def change(name,data):
            if name=='xl/_rels/workbook.xml.rels':
                root=ET.fromstring(data)
                for r in root:
                    if r.get('Type','').endswith('/worksheet'):r.set('TargetMode','External');r.set('Target','https://example.invalid/file.xml')
                return ET.tostring(root)
            return data
        with self.assertRaises(ValidationError):validate_upload('external.xlsx',modified_archive(change))

    def test_wrong_sheet_root(self):
        data=modified_archive(lambda name,data:b'<dummy/>' if name=='xl/worksheets/sheet1.xml' else data)
        with self.assertRaises(ValidationError):validate_upload('sheet.xlsx',data)

    def test_content_type_mismatch(self):
        data=modified_archive(lambda name,data:data.replace(b'spreadsheetml.sheet.main+xml',b'not-an-excel-type') if name=='[Content_Types].xml' else data)
        with self.assertRaises(ValidationError):validate_upload('type.xlsx',data)

    def test_cell_limit(self):
        with patch('modules.validators.MAX_XLSX_CELLS',1):
            with self.assertRaises(ValidationError):validate_upload('cells.xlsx',synthetic_xlsx())
