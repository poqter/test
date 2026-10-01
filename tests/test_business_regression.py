"""Synthetic workbooks and fixed business outcomes for the four extracted modules.

Fixtures contain no real customer data. These are parser/model/export checks,
not a claim that every insurance-company source layout has been verified.
"""
from datetime import date
from io import BytesIO
import unittest
from unittest.mock import patch
from tests.streamlit_stub import install
st=install(strict_numeric=True)
from openpyxl import Workbook,load_workbook
import pandas as pd
from modules.performance import commission_calculator as commission
from modules.performance import summer
from modules.consultation import analyzer,remodeling


def xlsx_bytes(wb):
    buffer=BytesIO();wb.save(buffer);wb.close();return buffer.getvalue()


def coverage_fixture():
    wb=Workbook();contracts=wb.active;contracts.title='계약사항'
    coverage=wb.create_sheet('상품별보장내용')
    contracts['B2']='테스트고객님';contracts['D2']='35세'
    for col,val in ((10,24000000),(11,6000000),(12,18000000)):contracts.cell(9,col,val)
    for row,val in ((2,'테스트생명'),(3,'테스트 보장상품'),(4,'100세'),(5,'60/240'),(6,'월납'),(7,123456)):coverage.cell(row,6,val)
    coverage['B9']='일반암';coverage['F9']=30000000
    coverage['B10']='뇌혈관';coverage['F10']=10000000
    return xlsx_bytes(wb)


class BusinessRegression(unittest.TestCase):
    def setUp(self):st.session_state.clear()
    def test_commission_table_values_survive_split(self):
        wb=Workbook();ws=wb.active;ws.title='테스트생명'
        ws.append(['상품명','납입기간','1차년계','총수수료'])
        ws.append(['테스트 보장상품','20년',4.0,8.0]);ws.append(['끝','','',''])
        products,warnings=commission.parse_commission_workbook(xlsx_bytes(wb),'생명보험')
        self.assertEqual(len(products),1);self.assertFalse(warnings)
        self.assertEqual(products[0]['first_year_rate'],4.0)
        self.assertEqual(products[0]['total_rate'],8.0)
    def test_holding_parse_exact_premium(self):
        wb=Workbook();ws=wb.active
        ws.append(['증권번호','상품명','보험사','계약일','납입기간','납입기간구분','쉐어율','계약자','수금자명','계속보험료','계약상태'])
        ws.append(['TEST-001','테스트 종신보험','한화생명','2026-09-01',20,'년',100,'가상고객','가상설계사',123456,'정상'])
        results=commission.parse_holding_workbook(xlsx_bytes(wb))
        self.assertEqual(len(results),1)
        self.assertEqual(results[0]['premium'],123456)
        self.assertEqual(results[0]['contract_month'],'2026-09')
    def test_commission_export_uses_fixed_reference_totals(self):
        contracts=[{'premium':100000,'first_year_rate':4.0,'total_rate':8.0,'insurer':'테스트생명',
          'product':'테스트상품','conditions':'20년납','customer':'가상고객','policy_number':'TEST-001','collector':'가상설계사'}]
        payload=commission._make_excel(contracts,.65,'2026-09',[])
        wb=load_workbook(BytesIO(payload),data_only=False)
        try:
            self.assertEqual(wb['수수료 계산']['B4'].value,260000)
            self.assertEqual(wb['수수료 계산']['D4'].value,520000)
            self.assertEqual(wb['수수료 계산']['E7'].value,100000)
        finally:wb.close()
    def test_analyzer_parser_and_xlsx_export(self):
        payload=coverage_fixture();parsed=analyzer.parse_source_file(payload)
        self.assertEqual(parsed['contracts'][0]['monthly'],123456)
        self.assertEqual(parsed['coverages'][0]['values'][0],30000000)
        output,filename,customer=analyzer.build_analysis_file(payload,['일반암','뇌혈관'])
        wb=load_workbook(BytesIO(output),data_only=False)
        try:
            self.assertIn('보장 분석',wb.sheetnames)
            vals=[cell.value for sheet in wb for row in sheet for cell in row]
            self.assertIn(30000000,vals)
            self.assertTrue(any(isinstance(v,str) and v.startswith('=') for v in vals))
        finally:wb.close()
    def test_analyzer_parser_cache_returns_private_copy(self):
        from modules.shared import runtime_cache
        state={};payload=coverage_fixture()
        with patch.object(runtime_cache,'_session',return_value=state):
            a=analyzer.parse_source_file(payload);a['contracts'][0]['monthly']=1
            b=analyzer.parse_source_file(payload)
            self.assertEqual(b['contracts'][0]['monthly'],123456)
            self.assertEqual(runtime_cache.statistics(state=state)['builds'],1)
    def test_summer_share_normalization_and_rates(self):
        data=pd.DataFrame([
          {'보험사':'한화생명','계속보험료':100000,'계약일자':'2026-07-01','납입기간':'20년','쉐어율':100,'상품명':'보장보험'},
          {'보험사':'DB손해보험','계속보험료':30000,'계약일자':'2026-07-01','납입기간':'20년','쉐어율':30,'상품명':'보장보험'},
          {'보험사':'삼성화재','계속보험료':100000,'계약일자':'2026-07-01','납입기간':'10년','쉐어율':100,'상품명':'보장보험'}])
        result=summer.compute_summer(data)
        self.assertEqual(result['실적보험료'].tolist(),[100000,50000,100000])
        self.assertEqual(result['썸머율'].tolist(),[150,250,50])
        self.assertEqual(result['썸머환산금액'].tolist(),[150000,125000,50000])
    def test_summer_grade_boundaries_keep_existing_rules(self):
        self.assertEqual(summer.get_summer_grade(2999999)[0],'미달성')
        self.assertEqual(summer.get_summer_grade(3000000)[0],'일반')
        self.assertEqual(summer.get_summer_grade(15000000)[0],'크라운')
        self.assertEqual(summer.get_summer_grade(15000001)[0],'HWARANG')
    def test_remodeling_model_and_export(self):
        person=remodeling.Person('가상고객',200000,30000000,50000,6000000,[remodeling.NewPlan('신규안',100000,20)])
        self.assertEqual(person.after_monthly,150000)
        self.assertEqual(person.after_total,30000000)
        self.assertEqual(person.monthly_change,-50000)
        stream=remodeling.create_excel([person],'테스트 비교안',date(2026,10,1),'가상담당자')
        wb=load_workbook(stream,data_only=False)
        try:
            self.assertGreaterEqual(len(wb.sheetnames),1)
            self.assertIn('테스트 비교안',[c.value for ws in wb for row in ws for c in row])
        finally:wb.close()

if __name__=='__main__':unittest.main()
