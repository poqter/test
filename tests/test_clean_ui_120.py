from io import BytesIO
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest


def test_deposit_changed_inputs_hide_stale_results():
    app=AppTest.from_string('from modules.calculators.comparison.deposit_vs_shortpay import run\nrun()',default_timeout=30).run()
    assert not app.exception
    next(b for b in app.button if '10년 보유자산 비교' in b.label).click().run()
    assert not app.exception and app.session_state['hwarang_ds_result']['monthly']==100
    next(x for x in app.number_input if x.label=='공통 월납입금액 (만원)').set_value(200).run()
    assert any('다시 계산' in w.value for w in app.warning)
    next(b for b in app.button if '10년 보유자산 비교' in b.label).click().run()
    assert not app.exception and app.session_state['hwarang_ds_result']['monthly']==200
    assert not app.warning


def test_question_edit_preview_preserves_order():
    app=AppTest.from_string('from modules.consultation.question_studio import run\nrun()',default_timeout=30).run()
    assert not app.exception
    next(b for b in app.button if b.label=='추천 질문 10개 추가').click().run()
    before=list(app.session_state['b_questions'])
    next(b for b in app.button if b.label=='선택한 질문 편집·PDF 저장 →').click().run()
    assert not app.exception
    assert app.session_state['b_questions']==before
    assert any(x.label=='질문 내용 미리보기' for x in app.expander)
    assert len(app.get('download_button'))>=1


def test_pdf_export_preserves_unknown_payment_no_comments_and_death_order():
    from test_analyzer_kb_pdf import fixture
    from modules.consultation.analyzer_pdf import convert_kb
    from modules.consultation.analyzer_exports import build_analysis_from_data
    raw=fixture();raw['contracts'][0].update(premium_missing=True,monthly=0)
    data=convert_kb(raw,'2026-01')
    payload,_,_=build_analysis_from_data(data)
    book=load_workbook(BytesIO(payload))
    for sheet in book:
        assert not any(c.comment for row in sheet for c in row)
        death=[(c.row,str(c.value).replace(' ','')) for row in sheet for c in row if c.value and str(c.value).replace(' ','') in ('일반사망','질병사망')]
        positions={label:row for row,label in death}
        assert positions['일반사망']+1==positions['질병사망']
    assert all(book['보장 분석'].cell(row,4).value==0 for row in (8,9,10))
