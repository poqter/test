"""Run with: python -m unittest discover -s tests -v. Uses only synthetic data."""

import io
import unittest
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from streamlit.testing.v1 import AppTest


ROOT = Path(__file__).resolve().parents[1]
APP_IDS = (
    "analyzer", "remodeling", "deposit_vs_shortpay", "renewal_vs_nonrenewal",
    "inheritance_tax", "insurer_portal", "insurance_claim_guide",
    "silson_generation_comparison", "convention", "summer", "manager_results",
    "commission_calculator", "quick_calculators", "consultation_helper", "comparison_builder", "education_center",
)


def app_for(user="Admin", page="home"):
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=45)
    at.session_state["password_correct"] = True
    at.session_state["login_user"] = user
    at.session_state["active_app"] = page
    return at.run()


class WorkspaceTests(unittest.TestCase):
    def assert_clean(self, at):
        self.assertEqual([e.message for e in at.exception], [])

    def test_login_without_settings_fails_closed(self):
        at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=45).run()
        at.text_input[0].set_value("synthetic-password")
        at.button[0].click().run()
        self.assert_clean(at)
        self.assertFalse(at.session_state["password_correct"])
        self.assertTrue(at.error)

    def test_valid_login_and_logout(self):
        at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=45)
        at.secrets["passwords"] = {"Admin": "test-only-password"}
        at.run()
        at.text_input[0].set_value("test-only-password")
        at.button[0].click().run()
        self.assert_clean(at)
        self.assertEqual(at.session_state["login_user"], "Admin")
        at.button(key="v2_logout").click().run()
        self.assert_clean(at)
        self.assertFalse(at.session_state["password_correct"])
        self.assertIsNone(at.session_state["login_user"])

    def test_empty_unknown_and_wrong_passwords_rejected(self):
        for passwords, entered in (({"Admin": ""}, ""), ({"Unknown": "test"}, "test"), ({"Admin": "right"}, "wrong")):
            with self.subTest(passwords=passwords):
                at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=45)
                at.secrets["passwords"] = passwords
                at.run()
                at.text_input[0].set_value(entered)
                at.button[0].click().run()
                self.assert_clean(at)
                self.assertFalse(at.session_state["password_correct"])

    def test_all_pages_render(self):
        for app_id in APP_IDS:
            with self.subTest(app_id=app_id):
                at = app_for(page=app_id)
                self.assert_clean(at)
                self.assertFalse(at.error, [e.value for e in at.error])
                self.assertEqual(at.session_state["active_app"], app_id)
                at.button(key="v2_nav_home").click().run()
                self.assert_clean(at)
                self.assertEqual(at.session_state["active_app"], "home")

    def test_all_roles_home_render(self):
        for user in ("Admin", "Manager1", "Basic", "Crew", "Dream"):
            with self.subTest(user=user):
                at = app_for(user)
                self.assert_clean(at)
                self.assertFalse(at.error, [e.value for e in at.error])
                self.assertEqual(len([b for b in at.button if b.key.startswith("v2_launch_all_")]), len(at.session_state["ws_allowed_ids"]))

    def test_search_and_empty_result(self):
        at = app_for()
        at.text_input(key="v2_global_search").set_value("청구서류").run()
        self.assert_clean(at)
        self.assertEqual(len([b for b in at.button if b.key.startswith("v2_launch_search_")]), 1)
        at.button(key="v2_launch_search_insurance_claim_guide").click().run()
        self.assert_clean(at)
        self.assertEqual(at.session_state["active_app"], "insurance_claim_guide")
        at.button(key="v2_nav_home").click().run()
        at.text_input(key="v2_global_search").set_value("zz-no-match").run()
        self.assert_clean(at)
        self.assertTrue(at.info)

    def test_permissions_and_invalid_routes(self):
        for route in ("commission_calculator", "not-a-page"):
            at = app_for("Basic", route)
            self.assert_clean(at)
            self.assertEqual(at.session_state["active_app"], "home")
            self.assertNotIn("v2_launch_all_commission_calculator", [b.key for b in at.button])
        at = app_for("Basic")
        self.assertFalse(at.button(key="v2_launch_all_analyzer").disabled)

    def test_rapid_repeated_navigation(self):
        at = app_for()
        for _ in range(3):
            at.button(key="v2_launch_quick_analyzer").click().run()
            self.assert_clean(at)
            at.button(key="v2_nav_home").click().run()
            self.assert_clean(at)

    def test_restored_assets(self):
        self.assertTrue((ROOT / ".streamlit/config.toml").is_file())
        self.assertTrue((ROOT / "assets/fonts/PretendardVariable.ttf").is_file())
        self.assertGreaterEqual(len(list((ROOT / "assets/insurer_logos").glob("*.png"))), 20)

    def test_summer_excel_missing_values(self):
        from modules.summer import excel_safe_value
        wb = Workbook()
        for index, value in enumerate((pd.NA, pd.NaT, np.nan, None, np.int64(3), "테스트"), start=1):
            wb.active.cell(index, 1, excel_safe_value(value))
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        reopened = load_workbook(out)
        self.assertIsNone(reopened.active["A1"].value)
        self.assertEqual(reopened.active["A5"].value, 3)
        self.assertEqual(reopened.active["A6"].value, "테스트")

    def test_analysis_excel_round_trip(self):
        from modules.analyzer import build_analysis_file
        source = Workbook()
        contracts = source.active
        contracts.title = "계약사항"
        contracts["B2"] = "가상고객"
        contracts["D2"] = "40세"
        contracts["J9"], contracts["K9"], contracts["L9"] = 24000000, 12000000, 12000000
        coverage = source.create_sheet("상품별보장내용")
        for row, value in enumerate(("가상보험", "가상상품", "100세", "120/240", "월납", 100000), start=2):
            coverage.cell(row, 6, value)
        coverage["B9"], coverage["F9"] = "일반암", 5000
        out = io.BytesIO()
        source.save(out)
        content, filename, _ = build_analysis_file(out.getvalue(), selected_labels=["일반암"])
        result = load_workbook(io.BytesIO(content))
        self.assertEqual(result.sheetnames, ["보장 분석", "보장 제안서"])
        self.assertEqual(result["보장 분석"]["D7"].value, 100000)
        self.assertEqual(result["보장 분석"]["D12"].value, 5000)
        self.assertEqual(result["보장 분석"]["A1"].font.sz, 14)
        self.assertIn("40세", result["보장 제안서"]["A1"].value)
        self.assertTrue(filename.endswith(".xlsx"))


if __name__ == "__main__":
    unittest.main()
