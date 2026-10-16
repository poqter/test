"""Offline regression tests; customer files are never bundled with the patch."""
import copy
from io import BytesIO
import unittest
from openpyxl import load_workbook
from modules.consultation.analyzer_pdf import amount, classify, convert_kb, extract_kb_pdf
from modules.consultation.analyzer_pdf_ui import apply_review
from modules.consultation.analyzer_exports import build_analysis_from_data


def fixture():
    return dict(customer_name='테스트',age='35',source_date='2026-01-29',contracts=[dict(
        company='테스트보험',product='합성보험',start='2025-12-29',end='2080-12-29',cycle='월납',
        years=20,monthly=10000,premium_missing=False,fully_paid=False,pages=[7],rows=[
            dict(number=1,kind='정액',raw='암진단비',standard='암진단',value=1000,page=7)])])


class TestKBPDF(unittest.TestCase):
    def test_amount_units(self):
        self.assertEqual(amount('1억 2,000만'),12000)
        self.assertEqual(amount('5,000원'),.5)
        self.assertEqual(amount('3.5만'),3.5)
        self.assertEqual(amount('0'),0)
        with self.assertRaises(ValueError):amount('알수없음')

    def test_disability_threshold(self):
        targets,reason=classify(dict(kind='정액',raw='상해후유장해(20~100%)',standard='상해후유장해',value=100))
        self.assertNotIn('재해(상해)장해3%일경우',dict(targets))
        self.assertIn('제외',reason)

    def test_targeted_drug_not_radiation(self):
        base=dict(kind='정액',standard='고액항암치료비',value=1000)
        self.assertTrue(classify(dict(base,raw='표적항암약물허가치료비'))[0])
        self.assertFalse(classify(dict(base,raw='표적항암방사선치료비'))[0])

    def test_zoster_eye_not_general(self):
        self.assertFalse(classify(dict(kind='정액',standard='특정질병진단',raw='대상포진눈병진단비',value=100))[0])

    def test_month_count_and_bound(self):
        raw=fixture();data=convert_kb(raw,'2026-01')
        self.assertEqual(data['contracts'][0]['payment_count'],'2/240회')
        self.assertEqual(data['contracts'][0]['paid'],20000)
        self.assertEqual(convert_kb(raw,'2050-01')['contracts'][0]['payment_count'],'240/240회')
        with self.assertRaises(ValueError):convert_kb(raw,'2025-12')

    def test_unknown_payment(self):
        raw=fixture();raw['contracts'][0]['premium_missing']=True;raw['contracts'][0]['monthly']=0
        self.assertTrue(convert_kb(raw,'2026-01')['contracts'][0]['payment_unknown'])

    def test_combined_limit_not_double_sum(self):
        raw=fixture();r=raw['contracts'][0]['rows'][0]
        r.update(kind='실손',standard='질병(전체질병을 의미)의료비(입원+통원)')
        data=convert_kb(raw,'2026-01')
        vals={x['label']:x['values'][0] for x in data['coverages']}
        self.assertEqual(vals['질병입원(실손)'],'확인 필요')
        self.assertEqual(vals['질병통원(실손)'],'확인 필요')

    def test_unmapped_retained(self):
        raw=fixture();raw['contracts'][0]['rows'][0]['standard']='null'
        data=convert_kb(raw,'2026-01');self.assertEqual(len(data['details']),1)
        self.assertEqual(data['details'][0]['targets'],[])

    def test_bad_pdf_rejected(self):
        with self.assertRaises(ValueError):extract_kb_pdf(b'not a pdf')

    def test_review_validation_and_export_gate(self):
        data=convert_kb(fixture(),'2026-01')
        with self.assertRaises(ValueError):build_analysis_from_data(data)
        rows=[{'1번':str(x['values'][0])} for x in data['coverages']]
        with self.assertRaises(ValueError):apply_review(data,'=formula',35,rows)
        bad=copy.deepcopy(rows);bad[0]['1번']='-1'
        with self.assertRaises(ValueError):apply_review(data,'테스트',35,bad)
        reviewed=apply_review(data,'테스트',35,rows)
        result,_,_=build_analysis_from_data(reviewed)
        w=load_workbook(BytesIO(result));self.assertEqual(w.sheetnames,['보장 분석','보장 제안서'])
        self.assertIn('35세',w['보장 분석']['A1'].value)
        self.assertIn('COUNTIF',w['보장 분석']['A7'].value)

    def test_review_does_not_mutate_parser(self):
        data=convert_kb(fixture(),'2026-01');rows=[{'1번':str(x['values'][0])} for x in data['coverages']]
        rows[0]['1번']='500'
        result=apply_review(data,'테스트',35,rows)
        self.assertEqual(data['coverages'][0]['values'][0],0)
        self.assertEqual(result['coverages'][0]['values'][0],500)


if __name__=='__main__':unittest.main()
